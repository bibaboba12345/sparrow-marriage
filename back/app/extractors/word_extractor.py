"""DOCX (Word) text extraction via python-docx."""

from __future__ import annotations

import io
from typing import Any

from .clean import clean_text


class WordExtractError(Exception):
    pass


def extract_docx(data: bytes, *, filename: str = "document.docx") -> dict[str, Any]:
    if not data:
        raise WordExtractError("Empty DOCX")

    lower = filename.lower()
    if lower.endswith(".doc") and not lower.endswith(".docx"):
        raise WordExtractError(
            "Legacy .doc not supported — save as .docx or paste text manually"
        )

    try:
        from docx import Document
    except ImportError as exc:
        raise WordExtractError("python-docx is not installed") from exc

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise WordExtractError(f"Failed to open DOCX: {exc}") from exc

    parts: list[str] = []

    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)

    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))

    text = clean_text("\n".join(parts))
    if not text:
        raise WordExtractError(f"No text found in {filename}")

    return {
        "filename": filename,
        "filetype": "docx",
        "engine": "python-docx",
        "page_count": None,
        "char_count": len(text),
        "text": text,
        "warnings": [],
    }
